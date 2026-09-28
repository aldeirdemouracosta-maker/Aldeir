"""Exportação PowerPoint (.pptx) com python-pptx."""
from __future__ import annotations

import base64
import io
import re
from typing import Any

SLIDE_W_EMU = 12192000  # 13.333in
SLIDE_H_EMU = 6858000   # 7.5in
UNIT_EMU = SLIDE_W_EMU / 100.0  # 1 unidade do canvas (a altura usa a mesma escala)


def _imports():
    try:
        from pptx import Presentation
        from pptx.dml.color import RGBColor
        from pptx.enum.shapes import MSO_SHAPE
        from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE, PP_ALIGN
        from pptx.oxml.ns import qn
        from pptx.util import Emu, Pt
    except ImportError as exc:
        raise RuntimeError("Exportação PPTX requer python-pptx. Instale com: pip install -r requirements.txt") from exc
    return Presentation, RGBColor, MSO_SHAPE, MSO_ANCHOR, MSO_AUTO_SIZE, PP_ALIGN, qn, Emu, Pt


def _hex(value: Any, fallback: str) -> str:
    if isinstance(value, str) and re.fullmatch(r"#[0-9A-Fa-f]{6}", value):
        return value[1:].upper()
    return fallback


def _rgb(value: Any, fallback: str = "FFFFFF"):
    RGBColor = _imports()[1]
    return RGBColor.from_string(_hex(value, fallback))


def _emu(v: Any) -> int:
    try:
        return int(round(float(v) * UNIT_EMU))
    except (TypeError, ValueError):
        return 0


def _decode_image(data_url: str) -> bytes | None:
    m = re.fullmatch(r"data:image/(?:png|jpeg|jpg|gif|bmp);base64,(.+)", data_url or "", re.S)
    if not m:
        return None
    try:
        return base64.b64decode(m.group(1), validate=True)
    except Exception:
        return None


def _set_alpha(shape, opacity: float) -> None:
    qn = _imports()[6]
    if opacity >= 0.999:
        return
    fill = shape._element.spPr.find(qn("a:solidFill"))
    if fill is None or not len(fill):
        return
    color = fill[0]
    for old in color.findall(qn("a:alpha")):
        color.remove(old)
    alpha = color.makeelement(qn("a:alpha"), {"val": str(int(max(0.0, opacity) * 100000))})
    color.append(alpha)


def _add_bullet(paragraph, font_pt: float) -> None:
    _, _, _, _, _, _, qn, _, Pt = _imports()
    ppr = paragraph._p.get_or_add_pPr()
    indent = int(Pt(font_pt * 1.1))
    ppr.set("marL", str(indent))
    ppr.set("indent", str(-indent))
    for tag in ("a:buNone", "a:buChar", "a:buAutoNum"):
        for old in ppr.findall(qn(tag)):
            ppr.remove(old)
    ppr.append(ppr.makeelement(qn("a:buChar"), {"char": "•"}))


def _crop_cover(pic, raw: bytes, box_w: int, box_h: int) -> None:
    """Recorta a imagem para preencher a caixa sem distorcer (object-fit: cover)."""
    try:
        from PIL import Image
        with Image.open(io.BytesIO(raw)) as img:
            iw, ih = img.size
    except Exception:
        return
    if not iw or not ih or not box_w or not box_h:
        return
    img_ratio, box_ratio = iw / ih, box_w / box_h
    if abs(img_ratio - box_ratio) < 1e-3:
        return
    if img_ratio > box_ratio:
        cut = (1 - box_ratio / img_ratio) / 2
        pic.crop_left = pic.crop_right = cut
    else:
        cut = (1 - img_ratio / box_ratio) / 2
        pic.crop_top = pic.crop_bottom = cut


def build_pptx(presentation: dict[str, Any]) -> bytes:
    Presentation, _, MSO_SHAPE, MSO_ANCHOR, MSO_AUTO_SIZE, PP_ALIGN, _, Emu, Pt = _imports()
    prs = Presentation()
    prs.slide_width = Emu(SLIDE_W_EMU)
    prs.slide_height = Emu(SLIDE_H_EMU)
    blank = prs.slide_layouts[6]
    theme = presentation.get("theme", {})
    fg = theme.get("foreground", "#16181D")

    for slide_data in presentation.get("slides", []):
        slide = prs.slides.add_slide(blank)
        bg = slide.background.fill
        bg.solid()
        bg.fore_color.rgb = _rgb(slide_data.get("background") or theme.get("background", "#FFFFFF"), "FFFFFF")
        credits: list[str] = []

        for el in slide_data.get("elements", []):
            etype = el.get("type")
            left, top, width, height = _emu(el.get("x", 0)), _emu(el.get("y", 0)), _emu(el.get("w", 20)), _emu(el.get("h", 10))
            width, height = max(width, 1), max(height, 1)
            if etype == "text":
                box = slide.shapes.add_textbox(left, top, width, height)
                tf = box.text_frame
                tf.clear()
                tf.word_wrap = True
                tf.auto_size = MSO_AUTO_SIZE.NONE
                tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
                tf.vertical_anchor = {"middle": MSO_ANCHOR.MIDDLE, "bottom": MSO_ANCHOR.BOTTOM}.get(el.get("valign"), MSO_ANCHOR.TOP)
                size = max(6.0, min(200.0, float(el.get("fontSize", 20) or 20)))
                bullets = bool(el.get("bullets"))
                lines = str(el.get("text", "")).replace("\r\n", "\n").replace("\r", "\n").split("\n")
                for i, line in enumerate(lines or [""]):
                    p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
                    p.alignment = {"center": PP_ALIGN.CENTER, "right": PP_ALIGN.RIGHT}.get(el.get("align"), PP_ALIGN.LEFT)
                    p.space_after = Pt(size * (0.45 if bullets else 0.2))
                    run = p.add_run()
                    run.text = line
                    font = run.font
                    font.size = Pt(size)
                    font.bold = bool(el.get("bold", False))
                    font.italic = bool(el.get("italic", False))
                    font.color.rgb = _rgb(el.get("color", fg), _hex(fg, "16181D"))
                    if el.get("font"):
                        font.name = str(el["font"])[:60]
                    if bullets and line.strip():
                        _add_bullet(p, size)
            elif etype == "shape":
                kind = el.get("shape")
                radius = float(el.get("radius", 0) or 0)
                if kind == "ellipse":
                    shape_type = MSO_SHAPE.OVAL
                elif radius > 0:
                    shape_type = MSO_SHAPE.ROUNDED_RECTANGLE
                else:
                    shape_type = MSO_SHAPE.RECTANGLE
                shape = slide.shapes.add_shape(shape_type, left, top, width, height)
                if shape_type == MSO_SHAPE.ROUNDED_RECTANGLE:
                    short = max(0.01, min(float(el.get("w", 1)), float(el.get("h", 1))))
                    shape.adjustments[0] = max(0.0, min(0.5, radius / short))
                shape.fill.solid()
                shape.fill.fore_color.rgb = _rgb(el.get("fill", theme.get("accent", "#6657FF")), "6657FF")
                _set_alpha(shape, float(el.get("opacity", 1) if el.get("opacity") is not None else 1))
                if el.get("stroke"):
                    shape.line.color.rgb = _rgb(el.get("stroke"), "333333")
                    shape.line.width = Pt(1.5)
                else:
                    shape.line.fill.background()
                shape.shadow.inherit = False
                if shape.has_text_frame:
                    shape.text_frame.text = ""
            elif etype == "image":
                raw = _decode_image(str(el.get("src", "")))
                if not raw:
                    continue
                pic = slide.shapes.add_picture(io.BytesIO(raw), left, top, width=width, height=height)
                if el.get("fit", "cover") == "cover":
                    _crop_cover(pic, raw, width, height)
                try:
                    pic._element.nvPicPr.cNvPr.set("descr", str(el.get("alt", "Imagem"))[:500])
                except Exception:
                    pass
                if el.get("credit"):
                    credits.append(str(el["credit"]))

        notes = str(slide_data.get("notes") or "").strip()
        if credits:
            notes = (notes + "\n\n" if notes else "") + "Créditos: " + " | ".join(credits)
        if notes:
            slide.notes_slide.notes_text_frame.text = notes

    props = prs.core_properties
    props.title = str(presentation.get("meta", {}).get("title", "MagicSlides"))[:250]
    props.author = str(presentation.get("meta", {}).get("author", ""))[:250]
    props.comments = "Criado com MagicSlides"
    out = io.BytesIO()
    prs.save(out)
    return out.getvalue()
