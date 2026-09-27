import base64
import io
import json
import threading
import urllib.error
import urllib.request
import zipfile

import pytest
from PIL import Image

from magicslides import ai, config, images, layouts, server
from magicslides.core import default_presentation, normalize_presentation, validate_presentation
from magicslides.html_export import build_html
from magicslides.pptx_export import build_pptx


def _png_bytes(w=64, h=40, color=(200, 80, 40)):
    out = io.BytesIO()
    Image.new("RGB", (w, h), color).save(out, "PNG")
    return out.getvalue()


def _data_url(w=64, h=40):
    return "data:image/png;base64," + base64.b64encode(_png_bytes(w, h)).decode()


@pytest.fixture(autouse=True)
def isolated_config(tmp_path, monkeypatch):
    monkeypatch.setenv("MAGICSLIDES_HOME", str(tmp_path))
    for env in config.ENV_FALLBACK.values():
        monkeypatch.delenv(env, raising=False)


# ------------------------------------------------------------------- core ---

def test_default_presentation_is_valid():
    p = default_presentation()
    assert p["schema"] == "magicslides.presentation"
    assert validate_presentation(p) == []


def test_legacy_gamma_livre_project_is_migrated():
    legacy = {"schema": "gamma-livre.presentation", "version": "0.1.0", "meta": {"title": "Antigo"},
              "theme": {"background": "#F7F7FA", "foreground": "#16181D", "accent": "#6657FF"},
              "slides": [{"id": "s1", "elements": [{"id": "t1", "type": "text", "x": 1, "y": 1, "w": 10, "h": 5, "text": "oi"}]}]}
    data = normalize_presentation(legacy)
    assert data["schema"] == "magicslides.presentation"
    assert validate_presentation(data) == []


def test_validation_rejects_remote_image_src_and_duplicates():
    p = default_presentation()
    el = p["slides"][0]["elements"][0]
    p["slides"][0]["elements"].append(dict(el))
    p["slides"][0]["elements"].append({"id": "img", "type": "image", "x": 0, "y": 0, "w": 10, "h": 10, "src": "http://x/y.png"})
    errors = validate_presentation(p)
    assert any("duplicado" in e for e in errors)
    assert any("data URL" in e for e in errors)


# ---------------------------------------------------------------- layouts ---

def test_every_layout_builds_valid_slides():
    specs = [{"layout": name, "title": f"Título {name}", "bullets": ["Um", "Dois", "Três"], "body": "Texto",
              "quote": "Uma frase", "quote_author": "Alguém", "image_query": "solar panels",
              "stats": [{"value": "73%", "label": "algo"}], "columns": [{"heading": "A", "bullets": ["x"]}, {"heading": "B", "bullets": ["y"]}]}
             for name in layouts.LAYOUTS]
    outline = layouts.normalize_outline({"title": "Teste", "slides": specs})
    pres = layouts.build_presentation(outline, "minimal")
    assert validate_presentation(pres) == []
    assert len(pres["slides"]) == len(layouts.LAYOUTS)
    kinds = {e["type"] for s in pres["slides"] for e in s["elements"]}
    assert kinds == {"text", "shape", "image"}


def test_normalize_outline_is_defensive():
    outline = layouts.normalize_outline({"slides": [{"layout": "???", "bullets": "a\nb", "stats": "x"}, 5, None]})
    assert outline["slides"][0]["layout"] == "title"
    assert outline["slides"][1]["layout"] == "bullets"
    assert outline["slides"][1]["bullets"] == ["a", "b"]


def test_fit_font_shrinks_long_text():
    short = layouts.fit_font("Oi", 40, 10, 40)
    long = layouts.fit_font("palavra " * 120, 40, 10, 40)
    assert short == 40 and long < 20


# --------------------------------------------------------------------- IA ---

def test_offline_skeleton_for_short_topic():
    outline, used = ai.generate_outline("Energia solar", 8, "pt", "", {"provider": "offline"})
    assert "offline" in used.lower()
    assert len(outline["slides"]) == 8
    assert outline["slides"][0]["layout"] == "title"
    assert outline["slides"][-1]["layout"] == "closing"


def test_offline_structures_pasted_text():
    text = "Minha palestra\n# Introdução\n- ponto um\n- ponto dois\n# Conclusão\nFrase um. Frase dois."
    outline, _ = ai.generate_outline(text, 10, "pt", "", {"provider": "offline"})
    titles = [s["title"] for s in outline["slides"]]
    assert titles[0] == "Minha palestra"
    assert "Introdução" in titles and "Conclusão" in titles
    intro = outline["slides"][titles.index("Introdução")]
    assert intro["bullets"] == ["ponto um", "ponto dois"]


def test_parse_json_object_handles_fences_and_noise():
    assert ai.parse_json_object('```json\n{"a": 1}\n```') == {"a": 1}
    assert ai.parse_json_object('Aqui está: {"a": 2} fim') == {"a": 2}
    with pytest.raises(ai.AIError):
        ai.parse_json_object("sem json")


def test_outline_schema_is_strict():
    def check(node):
        if node.get("type") == "object":
            assert node["additionalProperties"] is False
            assert set(node["required"]) == set(node["properties"])
            for child in node["properties"].values():
                check(child)
        if node.get("type") == "array":
            check(node["items"])
    check(ai.OUTLINE_SCHEMA)


def test_openai_compatible_provider(monkeypatch):
    captured = {}

    class FakeResp:
        def __init__(self, body):
            self.body = body

        def read(self):
            return self.body

        def __enter__(self):
            return self

        def __exit__(self, *a):
            return False

    def fake_urlopen(req, timeout=0):
        captured["url"] = req.full_url
        captured["body"] = json.loads(req.data)
        content = json.dumps({"title": "T", "subtitle": "", "slides": [{"layout": "bullets", "title": "A", "bullets": ["x"]}]})
        return FakeResp(json.dumps({"choices": [{"message": {"content": content}}]}).encode())

    monkeypatch.setattr(ai.urllib.request, "urlopen", fake_urlopen)
    cfg = {"provider": "openai", "openai_base_url": "http://localhost:11434/v1/", "openai_model": "llama3.1"}
    outline, used = ai.generate_outline("tema", 5, "pt", "", cfg)
    assert captured["url"] == "http://localhost:11434/v1/chat/completions"
    assert captured["body"]["model"] == "llama3.1"
    assert outline["slides"][1]["title"] == "A"
    assert "llama3.1" in used


def test_anthropic_provider_builds_structured_request(monkeypatch):
    anthropic = pytest.importorskip("anthropic")
    calls = {}

    class Block:
        type = "text"
        text = json.dumps({"title": "IA", "subtitle": "", "slides": [{"layout": "title", "title": "IA"}]})

    class Msg:
        stop_reason = "end_turn"
        content = [Block()]

    class Stream:
        def __enter__(self):
            return self

        def __exit__(self, *a):
            return False

        def get_final_message(self):
            return Msg()

    class FakeBetaMessages:
        def stream(self, **kw):
            calls.update(kw)
            return Stream()

    class FakeClient:
        def __init__(self, api_key=None):
            calls["api_key"] = api_key
            self.beta = type("B", (), {"messages": FakeBetaMessages()})()
            self.messages = FakeBetaMessages()

    monkeypatch.setattr(anthropic, "Anthropic", FakeClient)
    cfg = {"provider": "anthropic", "anthropic_api_key": "sk-test", "anthropic_model": "claude-opus-5", "anthropic_effort": "medium"}
    outline, _ = ai.generate_outline("tema", 4, "pt", "", cfg)
    assert outline["title"] == "IA"
    assert calls["api_key"] == "sk-test"
    assert calls["model"] == "claude-opus-5"
    assert calls["fallbacks"] == "default"
    assert calls["output_config"]["format"]["type"] == "json_schema"
    assert "temperature" not in calls


# ---------------------------------------------------------------- imagens ---

def test_fill_images_uses_results_and_placeholders(monkeypatch):
    def fake_source(query, per_page, cfg):
        if "nada" in query:
            return []
        return [{"id": f"x:{query}", "source": "openverse", "title": query, "thumb": "", "url": f"https://img.example/{query}.png",
                 "width": 800, "height": 500, "author": "Fulano", "license": "CC BY 4.0", "page": "https://example.org"}]

    monkeypatch.setitem(images.SOURCES, "openverse", fake_source)
    monkeypatch.setattr(images, "download_data_url", lambda url, max_side=1600: (_data_url(), 64, 40))
    outline = layouts.normalize_outline({"title": "T", "slides": [
        {"layout": "title", "title": "T"},
        {"layout": "image-right", "title": "A", "image_query": "sol"},
        {"layout": "image-left", "title": "B", "image_query": "nada"},
    ]})
    pres = layouts.build_presentation(outline, "aurora")
    credits, warnings = images.fill_presentation_images(pres, {"image_sources": ["openverse"]})
    assert len(credits) == 1 and "Fulano" in credits[0]
    imgs = [e for s in pres["slides"] for e in s["elements"] if e["type"] == "image"]
    assert len(imgs) == 1 and imgs[0]["src"].startswith("data:image/")
    placeholders = [e for s in pres["slides"] for e in s["elements"] if "placeholderQuery" in e]
    assert placeholders and placeholders[0]["placeholderQuery"] == "nada"
    assert validate_presentation(pres) == []


def test_ssrf_guard_blocks_local_addresses():
    for url in ("http://127.0.0.1/x.png", "http://localhost/x.png", "file:///etc/passwd", "http://10.0.0.1/a.jpg"):
        with pytest.raises(images.ImageError):
            images._check_public_url(url)


def test_redirect_to_private_address_is_blocked():
    handler = images._SafeRedirect()
    req = urllib.request.Request("https://example.org/a.png")
    with pytest.raises(images.ImageError):
        handler.redirect_request(req, None, 302, "Found", {}, "http://127.0.0.1/secret")


def test_encode_image_downscales_and_converts():
    raw = _png_bytes(3000, 1500)
    data_url, w, h = images.encode_image(raw, 1600)
    assert data_url.startswith("data:image/jpeg;base64,")
    assert (w, h) == (1600, 800)


def test_configured_sources_prefer_keyed_sources():
    cfg = {"image_sources": ["openverse"], "pexels_api_key": "k"}
    assert images.configured_sources(cfg) == ["pexels", "openverse"]
    assert images.configured_sources({"image_sources": ["pexels"]}) == []


# ----------------------------------------------------------- configuração ---

def test_settings_never_expose_keys():
    config.save({"provider": "anthropic", "anthropic_api_key": "sk-secret"})
    view = config.public_view()
    assert view["has_anthropic_api_key"] is True
    assert "sk-secret" not in json.dumps(view)
    config.save({"anthropic_api_key": ""})  # vazio mantém a chave
    assert config.load()["anthropic_api_key"] == "sk-secret"
    config.save({"anthropic_api_key": "__clear__"})
    assert config.load()["anthropic_api_key"] == ""


# ------------------------------------------------------------- exportação ---

def _rich_presentation():
    outline = layouts.normalize_outline({"title": "Export", "slides": [
        {"layout": "title", "title": "Export", "subtitle": "sub"},
        {"layout": "image-right", "title": "Com imagem", "bullets": ["a", "b"], "image_query": "q", "notes": "fale isso"},
        {"layout": "stats", "title": "Números", "stats": [{"value": "10%", "label": "x"}]},
        {"layout": "image-full", "title": "Full", "body": "texto", "image_query": "q2"},
    ]})
    pres = layouts.build_presentation(outline, "terracota")
    for s in pres["slides"]:
        for e in s["elements"]:
            if e["type"] == "image":
                e["src"] = _data_url(300, 100)
                e["credit"] = "Foto por Fulano"
    return pres


def test_pptx_export_contains_slides_notes_bullets_and_crop():
    from pptx import Presentation

    blob = build_pptx(_rich_presentation())
    assert blob[:2] == b"PK"
    prs = Presentation(io.BytesIO(blob))
    assert len(prs.slides) == 4
    assert prs.slide_width == 12192000
    s2 = prs.slides[1]
    assert "fale isso" in s2.notes_slide.notes_text_frame.text
    assert "Fulano" in s2.notes_slide.notes_text_frame.text
    pics = [sh for sh in s2.shapes if sh.shape_type == 13]
    assert pics and (pics[0].crop_left > 0 or pics[0].crop_top > 0)
    xml = zipfile.ZipFile(io.BytesIO(blob)).read("ppt/slides/slide2.xml").decode()
    assert "buChar" in xml
    xml4 = zipfile.ZipFile(io.BytesIO(blob)).read("ppt/slides/slide4.xml").decode()
    assert "<a:alpha" in xml4


def test_html_export_is_self_contained_and_escaped():
    pres = _rich_presentation()
    pres["slides"][0]["elements"][2]["text"] = "<script>alert(1)</script>"
    html = build_html(pres).decode()
    assert html.count('class="slide"') == 4
    assert "<script>alert(1)</script>" not in html
    assert "data:image/png;base64" in html


# ----------------------------------------------------------------- servidor ---

@pytest.fixture()
def live_server():
    httpd = server.make_server("127.0.0.1", 0)
    t = threading.Thread(target=httpd.serve_forever, daemon=True)
    t.start()
    yield f"http://127.0.0.1:{httpd.server_address[1]}"
    httpd.shutdown()
    httpd.server_close()


def _req(url, body=None, token=True, headers=None):
    h = {"Content-Type": "application/json", **(headers or {})}
    if token:
        h["X-MagicSlides-Token"] = server.TOKEN
    data = json.dumps(body).encode() if body is not None else None
    return urllib.request.urlopen(urllib.request.Request(url, data=data, headers=h), timeout=30)


def test_server_requires_token_and_injects_it(live_server):
    page = urllib.request.urlopen(live_server + "/").read().decode()
    assert server.TOKEN in page and "__MS_TOKEN__" not in page
    with pytest.raises(urllib.error.HTTPError) as err:
        _req(live_server + "/api/export/pptx", default_presentation(), token=False)
    assert err.value.code == 403
    with pytest.raises(urllib.error.HTTPError) as err:
        _req(live_server + "/api/settings", token=True, headers={"Host": "evil.example"})
    assert err.value.code == 403


def test_server_blocks_path_traversal(live_server):
    with pytest.raises(urllib.error.HTTPError) as err:
        urllib.request.urlopen(live_server + "/../server.py")
    assert err.value.code == 404


def test_server_generation_job_offline_without_images(live_server):
    import time

    job = json.load(_req(live_server + "/api/generate", {"prompt": "Reciclagem", "slides": 6, "images": False, "theme": "floresta"}))["job"]
    for _ in range(100):
        st = json.load(_req(f"{live_server}/api/jobs/{job}"))
        if st["state"] != "running":
            break
        time.sleep(0.05)
    assert st["state"] == "done", st
    pres = st["result"]
    assert len(pres["slides"]) == 6
    assert pres["theme"]["id"] == "floresta"
    assert validate_presentation(pres) == []
    assert not any(e["type"] == "image" for s in pres["slides"] for e in s["elements"])
    blob = _req(live_server + "/api/export/pptx", pres).read()
    assert blob[:2] == b"PK"


def test_server_validate_returns_migrated_presentation(live_server):
    legacy = default_presentation()
    legacy["schema"] = "gamma-livre.presentation"
    res = json.load(_req(live_server + "/api/validate", legacy))
    assert res["presentation"]["schema"] == "magicslides.presentation"


def test_server_encode_endpoint(live_server):
    res = json.load(_req(live_server + "/api/images/encode", {"dataUrl": _data_url(50, 50)}))
    assert res["src"].startswith("data:image/")
